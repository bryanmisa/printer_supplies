import csv
import io
import json
from datetime import datetime, timedelta
from decimal import Decimal
from itertools import chain

from django.db import models
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.contrib.contenttypes.models import ContentType
from django.core.paginator import Paginator
from django.db.models import Count, Q, Sum
from django.http import HttpResponse, JsonResponse, HttpResponseRedirect
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.utils.html import format_html
from django.views.decorators.csrf import csrf_exempt
from django.views import View
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView

from rbac.decorators import require_permission, require_role, admin_required

from .forms import (
    UserCreationForm, SupplierForm, PrinterModelForm, PrinterForm,
    SupplyForm, SupplyTypeForm, DeliveryForm, DeliveryItemForm, InstallationForm, DisposalForm,
    DepartmentForm, LocationForm, CustodianForm
)
from .models import (
    Supply, Printer, PrinterModel, Supplier, SupplyType, Delivery, DeliveryItem,
    SupplyInstallation, AuditLog, User, Department, Location, Custodian
)


def log_audit(user, action, obj, changes=None):
    if obj:
        AuditLog.objects.create(
            user=user,
            action=action,
            content_type=ContentType.objects.get_for_model(obj),
            object_id=obj.pk,
            object_repr=str(obj),
            changes=changes or {}
        )


def login_view(request):
    from django.contrib.auth.forms import AuthenticationForm

    remembered_username = request.COOKIES.get('remembered_username')

    if request.user.is_authenticated:
        return redirect('dashboard')

    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            login(request, user)

            remember_me = request.POST.get('remember_me')
            response = redirect(request.GET.get('next') or 'dashboard')

            if remember_me:
                response.set_cookie('remembered_username', form.cleaned_data.get('username'), max_age=timedelta(days=14))
            else:
                response.delete_cookie('remembered_username')

            return response
    else:
        initial = {'username': remembered_username} if remembered_username else {}
        form = AuthenticationForm(initial=initial)

    response = render(request, 'registration/login.html', {'form': form, 'remembered_username': remembered_username})
    return response


def logout_view(request):
    logout(request)
    return redirect('login')


@login_required
def dashboard(request):
    total_supplies = Supply.objects.filter(is_active=True).count()
    total_printers = Printer.objects.filter(status='active').count()
    installed_supplies = SupplyInstallation.objects.filter(is_active=True).count()
    out_of_stock = Supply.objects.filter(current_stock=0, is_active=True).count()
    low_stock = Supply.objects.filter(
        current_stock__gt=0,
        current_stock__lte=models.F('low_stock_threshold'),
        is_active=True
    ).count()

    out_of_stock_alerts = Supply.objects.filter(current_stock=0, is_active=True)[:5]
    low_stock_alerts = Supply.objects.filter(
        current_stock__gt=0,
        current_stock__lte=models.F('low_stock_threshold'),
        is_active=True
    )[:5]

    last_30_days = timezone.now() - timedelta(days=30)
    consumption_data = SupplyInstallation.objects.filter(
        disposed_at__gte=last_30_days
    ).values('disposed_at__date').annotate(count=Count('id')).order_by('disposed_at__date')

    labels = []
    data = []
    for i in range(30, -1, -1):
        date = (timezone.now() - timedelta(days=i)).date()
        labels.append(date.strftime('%b %d'))
        count = next((item['count'] for item in consumption_data if item['disposed_at__date'] == date), 0)
        data.append(count)

    supply_models = PrinterModel.objects.all()[:10]

    context = {
        'total_supplies': total_supplies,
        'total_printers': total_printers,
        'installed_supplies': installed_supplies,
        'out_of_stock': out_of_stock,
        'low_stock': low_stock,
        'out_of_stock_alerts': out_of_stock_alerts,
        'low_stock_alerts': low_stock_alerts,
        'chart_labels': json.dumps(labels),
        'chart_data': json.dumps(data),
        'supply_models': supply_models,
    }
    return render(request, 'inventory/dashboard.html', context)


class SupplyListView(LoginRequiredMixin, ListView):
    model = Supply
    template_name = 'inventory/supply_list.html'
    paginate_by = 20
    context_object_name = 'supplies'

    def get_queryset(self):
        queryset = Supply.objects.all().order_by('name')
        query = self.request.GET.get('q')
        status = self.request.GET.get('status')
        printer_model = self.request.GET.get('printer_model')
        supplier = self.request.GET.get('supplier')
        is_active = self.request.GET.get('is_active')

        if is_active == 'active':
            queryset = queryset.filter(is_active=True)
        elif is_active == 'inactive':
            queryset = queryset.filter(is_active=False)

        if query:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(sku__icontains=query)
            )
        if status:
            if status == 'out_of_stock':
                queryset = queryset.filter(current_stock=0)
            elif status == 'low_stock':
                queryset = queryset.filter(
                    current_stock__gt=0,
                    current_stock__lte=models.F('low_stock_threshold')
                )
            elif status == 'normal':
                queryset = queryset.filter(
                    current_stock__gt=models.F('low_stock_threshold')
                )
            elif status == 'overstock':
                queryset = queryset.filter(
                    current_stock__gt=models.F('max_stock_threshold')
                )
        if printer_model:
            queryset = queryset.filter(printer_models__id=printer_model)
        if supplier:
            queryset = queryset.filter(suppliers__id=supplier)

        return queryset.select_related()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['printer_models'] = PrinterModel.objects.all()
        context['suppliers'] = Supplier.objects.filter(is_active=True)
        return context


class SupplyCreateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        form = SupplyForm()
        return render(request, 'inventory/supply_form.html', {
            'form': form,
            'departments': Department.objects.all(),
            'locations': Location.objects.all(),
        })

    def post(self, request):
        form = SupplyForm(request.POST, request.FILES)
        if form.is_valid():
            supply = form.save()
            messages.success(request, 'Supply created successfully.')
            log_audit(request.user, 'create', supply, {'action': 'created'})
            return redirect('supply_list')
        return render(request, 'inventory/supply_form.html', {
            'form': form,
            'departments': Department.objects.all(),
            'locations': Location.objects.all(),
        })


class SupplyUpdateView(LoginRequiredMixin, View):
    model = Supply
    form_class = SupplyForm

    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        supply = get_object_or_404(Supply, pk=pk)
        form = SupplyForm(instance=supply)
        return render(request, 'inventory/supply_form.html', {
            'form': form,
            'object': supply,
            'departments': Department.objects.all(),
            'locations': Location.objects.all(),
        })

    def post(self, request, pk):
        supply = get_object_or_404(Supply, pk=pk)
        form = SupplyForm(request.POST, request.FILES, instance=supply)
        if form.is_valid():
            supply = form.save()
            messages.success(request, 'Supply updated successfully.')
            log_audit(request.user, 'update', supply, {'action': 'updated'})
            return redirect('supply_list')
        return render(request, 'inventory/supply_form.html', {
            'form': form,
            'object': supply,
            'departments': Department.objects.all(),
            'locations': Location.objects.all(),
        })


class SupplyDeleteView(LoginRequiredMixin, View):
    model = Supply
    template_name = 'inventory/supply_confirm_delete.html'

    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        supply = get_object_or_404(Supply, pk=pk)
        return render(request, 'inventory/supply_confirm_delete.html', {'supply': supply})

    def post(self, request, pk):
        supply = get_object_or_404(Supply, pk=pk)
        supply.is_active = False
        supply.save()
        messages.success(request, 'Supply deleted successfully.')
        log_audit(request.user, 'delete', supply, {'action': 'deleted'})
        return redirect('supply_list')


class SupplyDetailView(LoginRequiredMixin, DetailView):
    model = Supply
    template_name = 'inventory/supply_detail.html'
    context_object_name = 'supply'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        supply = self.object

        installations = SupplyInstallation.objects.filter(
            supply=supply
        ).select_related('printer', 'printer__printer_model', 'installed_by').order_by('-installed_at')

        context['installations'] = installations
        context['current_installations'] = installations.filter(is_active=True)
        context['installation_history'] = installations.filter(is_active=False)[:20]

        last_30_days = timezone.now() - timedelta(days=30)
        consumption_data = SupplyInstallation.objects.filter(
            supply=supply,
            disposed_at__gte=last_30_days,
            is_active=False
        ).values('disposed_at__date').annotate(count=Count('id')).order_by('disposed_at__date')

        labels = []
        data = []
        for i in range(30, -1, -1):
            date = (timezone.now() - timedelta(days=i)).date()
            labels.append(date.strftime('%b %d'))
            count = next((item['count'] for item in consumption_data if item['disposed_at__date'] == date), 0)
            data.append(count)

        context['chart_labels'] = json.dumps(labels)
        context['chart_data'] = json.dumps(data)

        return context


class SupplierListView(LoginRequiredMixin, ListView):
    model = Supplier
    template_name = 'inventory/supplier_list.html'
    paginate_by = 20
    context_object_name = 'suppliers'

    def get_queryset(self):
        queryset = Supplier.objects.all().order_by('name')
        query = self.request.GET.get('q')
        if query:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(contact_name__icontains=query)
            )
        return queryset


class SupplierDetailView(LoginRequiredMixin, DetailView):
    model = Supplier
    template_name = 'inventory/supplier_detail.html'
    context_object_name = 'supplier'


class SupplierCreateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        form = SupplierForm()
        if request.GET.get('modal'):
            return render(request, 'inventory/supplier_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/supplier_form.html', {'form': form})

    def post(self, request):
        form = SupplierForm(request.POST)
        if form.is_valid():
            supplier = form.save()
            messages.success(request, 'Supplier created successfully.')
            log_audit(request.user, 'create', supplier, {'action': 'created'})
            
            if request.POST.get('modal'):
                next_url = request.POST.get('next', reverse('supplier_list'))
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'pk': supplier.pk, 'name': supplier.name, 'next': next_url})
                return redirect(next_url + '?created=' + str(supplier.pk))
            return redirect('supplier_list')
        if request.POST.get('modal'):
            return render(request, 'inventory/supplier_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/supplier_form.html', {'form': form})


class SupplierUpdateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        supplier = get_object_or_404(Supplier, pk=pk)
        form = SupplierForm(instance=supplier)
        return render(request, 'inventory/supplier_form.html', {'form': form, 'object': supplier})

    def post(self, request, pk):
        supplier = get_object_or_404(Supplier, pk=pk)
        form = SupplierForm(request.POST, instance=supplier)
        if form.is_valid():
            supplier = form.save()
            messages.success(request, 'Supplier updated successfully.')
            log_audit(request.user, 'update', supplier, {'action': 'updated'})
            return redirect('supplier_list')
        return render(request, 'inventory/supplier_form.html', {'form': form, 'object': supplier})


class SupplierDeleteView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        supplier = get_object_or_404(Supplier, pk=pk)
        return render(request, 'inventory/supplier_confirm_delete.html', {'supplier': supplier})

    def post(self, request, pk):
        supplier = get_object_or_404(Supplier, pk=pk)
        supplier.is_active = False
        supplier.save()
        messages.success(request, 'Supplier deleted successfully.')
        log_audit(request.user, 'delete', supplier, {'action': 'deleted'})
        return redirect('supplier_list')


@login_required
def supplier_toggle_active(request, pk):
    supplier = get_object_or_404(Supplier, pk=pk)
    supplier.is_active = not supplier.is_active
    supplier.save()
    status = "activated" if supplier.is_active else "deactivated"
    messages.success(request, f'Supplier {status}.')
    return redirect('supplier_list')


class PrinterListView(LoginRequiredMixin, ListView):
    model = Printer
    template_name = 'inventory/printer_list.html'
    paginate_by = 20
    context_object_name = 'printers'

    def get_queryset(self):
        queryset = Printer.objects.all().order_by('name')
        query = self.request.GET.get('q')
        status = self.request.GET.get('status')
        custodian = self.request.GET.get('custodian')

        if query:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(serial_number__icontains=query)
            )
        if status:
            queryset = queryset.filter(status=status)
        if custodian:
            queryset = queryset.filter(custodian_id=custodian)

        return queryset.select_related('printer_model', 'custodian', 'department', 'location')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['custodians'] = Custodian.objects.all()
        return context


class PrinterCreateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        form = PrinterForm()
        return render(request, 'inventory/printer_form.html', {'form': form, 'departments': Department.objects.all(), 'locations': Location.objects.all()})

    def post(self, request):
        form = PrinterForm(request.POST, request.FILES)
        if form.is_valid():
            printer = form.save()
            messages.success(request, 'Printer created successfully.')
            log_audit(request.user, 'create', printer, {'action': 'created'})
            return redirect('printer_list')
        return render(request, 'inventory/printer_form.html', {'form': form, 'departments': Department.objects.all(), 'locations': Location.objects.all()})


class PrinterUpdateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        printer = get_object_or_404(Printer, pk=pk)
        form = PrinterForm(instance=printer)
        return render(request, 'inventory/printer_form.html', {'form': form, 'object': printer, 'departments': Department.objects.all(), 'locations': Location.objects.all()})

    def post(self, request, pk):
        printer = get_object_or_404(Printer, pk=pk)
        form = PrinterForm(request.POST, request.FILES, instance=printer)
        if form.is_valid():
            printer = form.save()
            messages.success(request, 'Printer updated successfully.')
            log_audit(request.user, 'update', printer, {'action': 'updated'})
            return redirect('printer_list')
        return render(request, 'inventory/printer_form.html', {'form': form, 'object': printer, 'departments': Department.objects.all(), 'locations': Location.objects.all()})


class PrinterDeleteView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        printer = get_object_or_404(Printer, pk=pk)
        return render(request, 'inventory/printer_confirm_delete.html', {'printer': printer})

    def post(self, request, pk):
        printer = get_object_or_404(Printer, pk=pk)
        messages.success(request, 'Printer deleted successfully.')
        log_audit(request.user, 'delete', printer, {'action': 'deleted'})
        printer.delete()
        return redirect('printer_list')


class PrinterDetailView(LoginRequiredMixin, DetailView):
    model = Printer
    template_name = 'inventory/printer_detail.html'
    context_object_name = 'printer'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        printer = self.object

        context['installations'] = SupplyInstallation.objects.filter(
            printer=printer
        ).select_related('supply', 'installed_by').order_by('-installed_at')

        context['current_installations'] = SupplyInstallation.objects.filter(
            printer=printer, is_active=True
        ).select_related('supply', 'installed_by')

        context['compatible_supplies'] = Supply.objects.filter(
            printer_models=printer.printer_model,
            is_active=True,
            current_stock__gt=0
        )

        return context


class PrinterModelListView(LoginRequiredMixin, ListView):
    model = PrinterModel
    template_name = 'inventory/printer_model_list.html'
    paginate_by = 20
    context_object_name = 'printer_models'

    def get_queryset(self):
        queryset = PrinterModel.objects.all().order_by('manufacturer', 'name')
        query = self.request.GET.get('q')
        manufacturer = self.request.GET.get('manufacturer')
        status = self.request.GET.get('status')

        if status != 'all':
            queryset = queryset.filter(is_active=True)

        if query:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(manufacturer__icontains=query)
            )
        if manufacturer:
            queryset = queryset.filter(manufacturer__icontains=manufacturer)

        return queryset.prefetch_related('printers', 'supplies')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['manufacturers'] = PrinterModel.objects.values_list('manufacturer', flat=True).distinct().order_by('manufacturer')
        return context


class PrinterModelDetailView(LoginRequiredMixin, DetailView):
    model = PrinterModel
    template_name = 'inventory/printer_model_detail.html'
    context_object_name = 'printer_model'


class PrinterModelCreateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        form = PrinterModelForm()
        if request.GET.get('modal'):
            return render(request, 'inventory/printer_model_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/printer_model_form.html', {'form': form})

    def post(self, request):
        form = PrinterModelForm(request.POST)
        if form.is_valid():
            printer_model = form.save()
            messages.success(request, 'Printer model created successfully.')
            
            if request.POST.get('modal'):
                next_url = request.POST.get('next', reverse('printer_model_list'))
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'pk': printer_model.pk, 'name': str(printer_model), 'next': next_url})
                return redirect(next_url + '?created=' + str(printer_model.pk))
            return redirect('printer_model_list')
        if request.POST.get('modal'):
            return render(request, 'inventory/printer_model_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/printer_model_form.html', {'form': form})


class PrinterModelUpdateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        model = get_object_or_404(PrinterModel, pk=pk)
        form = PrinterModelForm(instance=model)
        return render(request, 'inventory/printer_model_form.html', {'form': form, 'object': model})

    def post(self, request, pk):
        model = get_object_or_404(PrinterModel, pk=pk)
        form = PrinterModelForm(request.POST, instance=model)
        if form.is_valid():
            form.save()
            messages.success(request, 'Printer model updated successfully.')
            return redirect('printer_model_list')
        return render(request, 'inventory/printer_model_form.html', {'form': form, 'object': model})


class PrinterModelDeleteView(LoginRequiredMixin, View):
    def get(self, request, pk):
        model = get_object_or_404(PrinterModel, pk=pk)
        return render(request, 'inventory/printer_model_confirm_delete.html', {'printer_model': model})

    def post(self, request, pk):
        model = get_object_or_404(PrinterModel, pk=pk)
        model.printers.update(printer_model=None)
        log_audit(request.user, 'delete', model, {'action': 'deleted', 'printers_unlinked': True})
        model.delete()
        messages.success(request, 'Printer model deleted successfully.')
        return redirect('printer_model_list')


class PrinterModelToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        model = get_object_or_404(PrinterModel, pk=pk)
        model.is_active = not model.is_active
        model.save()
        status = 'enabled' if model.is_active else 'disabled'
        messages.success(request, f'Printer model {status} successfully.')
        return redirect('printer_model_list')


class DeliveryListView(LoginRequiredMixin, ListView):
    model = Delivery
    template_name = 'inventory/delivery_list.html'
    paginate_by = 20
    context_object_name = 'deliveries'

    def get_queryset(self):
        queryset = Delivery.objects.all().order_by('-delivery_date')
        query = self.request.GET.get('q')
        supplier = self.request.GET.get('supplier')
        
        if query:
            queryset = queryset.filter(notes__icontains=query)
        if supplier:
            queryset = queryset.filter(supplier_id=supplier)
        
        return queryset.select_related('supplier', 'created_by')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['suppliers'] = Supplier.objects.filter(is_active=True)
        return context


class DeliveryCreateView(LoginRequiredMixin, View):
    form_class = DeliveryForm
    template_name = 'inventory/delivery_form.html'

    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        form = self.form_class()
        supplier_id = request.GET.get('supplier')
        deliveries = None
        
        all_supplies = list(Supply.objects.filter(is_active=True).values('id', 'name', 'sku'))
        all_supplies_json = json.dumps(all_supplies)
        
        if supplier_id:
            deliveries = Delivery.objects.filter(supplier_id=supplier_id).select_related('supplier').prefetch_related('items', 'items__supply').order_by('-delivery_date')[:20]
            past_supplies = set()
            for d in deliveries:
                for item in d.items.all():
                    past_supplies.add(item.supply)
            context_supplies = list(past_supplies)
        else:
            context_supplies = []
        
        return render(request, self.template_name, {
            'form': form,
            'past_deliveries': deliveries,
            'past_supplies': context_supplies,
            'suppliers': Supplier.objects.all(),
            'all_supplies': all_supplies,
            'all_supplies_json': all_supplies_json,
            'today': datetime.now().strftime('%Y-%m-%d'),
            'departments': Department.objects.all(),
            'locations': Location.objects.all(),
        })

    def post(self, request):
        form = self.form_class(request.POST, request.FILES)
        if form.is_valid():
            delivery = form.save(commit=False)
            delivery.created_by = request.user
            delivery.save()

            supply_ids = request.POST.getlist('supplies')
            quantities = request.POST.getlist('quantities')

            for supply_id, quantity in zip(supply_ids, quantities):
                if supply_id and quantity:
                    supply = Supply.objects.get(pk=supply_id)
                    qty = int(quantity)
                    item, created = DeliveryItem.objects.get_or_create(
                        delivery=delivery,
                        supply=supply,
                        defaults={'quantity': qty}
                    )
                    if not created:
                        item.quantity += qty
                        item.save()
                    supply.current_stock += qty
                    supply.save()

                    log_audit(request.user, 'delivery', delivery, {
                        'supply': supply.name,
                        'quantity': str(qty)
                    })

            messages.success(request, 'Delivery created successfully.')
            return redirect('delivery_list')

        supplier_id = request.POST.get('supplier')
        deliveries = None

        all_supplies = list(Supply.objects.filter(is_active=True).values('id', 'name', 'sku'))
        all_supplies_json = json.dumps(all_supplies)

        if supplier_id:
            deliveries = Delivery.objects.filter(supplier_id=supplier_id).select_related('supplier').prefetch_related('items', 'items__supply').order_by('-delivery_date')[:20]
            past_supplies = set()
            for d in deliveries:
                for item in d.items.all():
                    past_supplies.add(item.supply)
            context_supplies = list(past_supplies)
        else:
            context_supplies = []

        return render(request, self.template_name, {
            'form': form,
            'past_deliveries': deliveries,
            'past_supplies': context_supplies,
            'suppliers': Supplier.objects.all(),
            'all_supplies': all_supplies,
            'all_supplies_json': all_supplies_json,
            'today': datetime.now().strftime('%Y-%m-%d'),
            'departments': Department.objects.all(),
            'locations': Location.objects.all(),
            'posted_supplier': request.POST.get('supplier'),
            'posted_supplies': request.POST.getlist('supplies'),
            'posted_quantities': request.POST.getlist('quantities'),
            'posted_notes': request.POST.get('notes'),
            'posted_delivery_date': request.POST.get('delivery_date'),
        })


class DeliveryDetailView(LoginRequiredMixin, DetailView):
    model = Delivery
    template_name = 'inventory/delivery_detail.html'
    context_object_name = 'delivery'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        delivery = self.object
        context['total_quantity'] = sum(item.quantity for item in delivery.items.all())
        return context


class InstallationListView(LoginRequiredMixin, ListView):
    model = SupplyInstallation
    template_name = 'inventory/installation_list.html'
    paginate_by = 20
    context_object_name = 'installations'

    def get_queryset(self):
        queryset = SupplyInstallation.objects.all().order_by('-installed_at')
        printer = self.request.GET.get('printer')
        status = self.request.GET.get('status')
        
        if printer:
            queryset = queryset.filter(printer_id=printer)
        if status == 'active':
            queryset = queryset.filter(is_active=True)
        elif status == 'disposed':
            queryset = queryset.filter(is_active=False)
        
        return queryset.select_related('printer', 'printer__printer_model', 'supply', 'installed_by')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['printers'] = Printer.objects.filter(status='active')
        return context


@login_required
def install_supply(request, pk):
    printer = get_object_or_404(Printer, pk=pk)

    if request.method == 'POST':
        supply_ids = request.POST.getlist('supplies')
        notes = request.POST.get('notes', '')

        if not supply_ids:
            messages.error(request, 'Please select at least one supply to install.')
            return redirect('printer_detail', pk=printer.id)

        installed = []
        errors = []

        for supply_id in supply_ids:
            supply = get_object_or_404(Supply, pk=supply_id)

            if supply.current_stock <= 0:
                errors.append(f'{supply.name} is out of stock.')
                continue

            if not supply.printer_models.filter(pk=printer.printer_model.pk).exists():
                errors.append(f'{supply.name} is not compatible with this printer model.')
                continue

            if SupplyInstallation.objects.filter(printer=printer, supply=supply, is_active=True).exists():
                errors.append(f'{supply.name} is already installed on this printer.')
                continue

            installation = SupplyInstallation.objects.create(
                printer=printer,
                supply=supply,
                installed_by=request.user,
                notes=notes
            )
            supply.current_stock -= 1
            supply.save()

            log_audit(request.user, 'install', installation, {
                'printer': printer.name,
                'supply': supply.name
            })
            installed.append(supply.name)

        if installed:
            messages.success(request, f'Installed: {", ".join(installed)}.')
        if errors:
            messages.error(request, f'Could not install: {", ".join(errors)}')

        return redirect('printer_detail', pk=printer.id)

    return redirect('printer_list')


@login_required
def dispose_supply(request, pk):
    installation = get_object_or_404(SupplyInstallation, pk=pk)
    printer_pk = installation.printer.pk
    
    if not installation.is_active:
        messages.error(request, 'This supply is already disposed.')
        return redirect('printer_detail', pk=printer_pk)

    if request.method == 'POST':
        notes = request.POST.get('notes', '')
        installation.is_active = False
        installation.disposed_at = timezone.now()
        installation.notes = notes
        installation.save()

        log_audit(request.user, 'dispose', installation, {
            'printer': installation.printer.name,
            'supply': installation.supply.name
        })

        messages.success(request, f'{installation.supply.name} disposed from {installation.printer.name}.')
        return redirect('printer_detail', pk=printer_pk)

    return redirect('printer_detail', pk=printer_pk)


@login_required
def inventory_report(request):
    supplies = Supply.objects.filter(is_active=True)
    
    out_of_stock = supplies.filter(current_stock=0)
    low_stock = supplies.filter(
        current_stock__gt=0,
        current_stock__lte=models.F('low_stock_threshold')
    )
    normal = supplies.filter(
        current_stock__gt=models.F('low_stock_threshold'),
        current_stock__lte=models.F('max_stock_threshold')
    )
    overstock = supplies.filter(
        current_stock__gt=models.F('max_stock_threshold')
    )

    fmt = request.GET.get('format')
    group = request.GET.get('group')
    
    if fmt and group:
        if group == 'out_of_stock':
            supplies = out_of_stock
        elif group == 'low_stock':
            supplies = low_stock
        elif group == 'normal':
            supplies = normal
        elif group == 'overstock':
            supplies = overstock
        else:
            supplies = Supply.objects.none()
        
        supplies = list(supplies)
        group_names = {'out_of_stock': 'Out of Stock', 'low_stock': 'Low Stock', 'normal': 'Normal Stock', 'overstock': 'Overstock'}
        
        if fmt == 'csv':
            return _export_group_csv(supplies, group_names.get(group, group), group)
        elif fmt == 'excel':
            return _export_group_excel(supplies, group_names.get(group, group), group)
        elif fmt == 'pdf':
            return _export_group_pdf(supplies, group_names.get(group, group), group)
    elif fmt:
        if fmt == 'csv':
            return _export_inventory_csv(out_of_stock, low_stock, normal, overstock)
        elif fmt == 'excel':
            return _export_inventory_excel(out_of_stock, low_stock, normal, overstock)
        elif fmt == 'pdf':
            return _export_inventory_pdf(out_of_stock, low_stock, normal, overstock)

    context = {
        'total_supplies': supplies.count(),
        'out_of_stock': out_of_stock,
        'out_of_stock_count': out_of_stock.count(),
        'low_stock': low_stock,
        'low_stock_count': low_stock.count(),
        'normal': normal,
        'normal_count': normal.count(),
        'overstock': overstock,
        'overstock_count': overstock.count(),
    }
    return render(request, 'inventory/inventory_report.html', context)


def _get_inventory_status(current, low_threshold, max_threshold):
    if current == 0:
        return 'Out of Stock'
    elif current <= low_threshold:
        return 'Low Stock'
    elif current > max_threshold:
        return 'Overstock'
    return 'Normal'


def _export_inventory_csv(out_of_stock, low_stock, normal, overstock):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="inventory_report.csv"'
    writer = csv.writer(response)
    
    writer.writerow(['== LOW STOCK =='])
    writer.writerow(['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status'])
    for supply in low_stock:
        writer.writerow([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, supply.max_stock_threshold, 'Low Stock'])
    
    writer.writerow([])
    writer.writerow(['== NORMAL STOCK =='])
    writer.writerow(['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status'])
    for supply in normal:
        writer.writerow([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, supply.max_stock_threshold, 'Normal'])
    
    writer.writerow([])
    writer.writerow(['== OVERSTOCK =='])
    writer.writerow(['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status'])
    for supply in overstock:
        writer.writerow([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, supply.max_stock_threshold, 'Overstock'])
    
    writer.writerow([])
    writer.writerow(['== OUT OF STOCK =='])
    writer.writerow(['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status'])
    for supply in out_of_stock:
        writer.writerow([supply.name, supply.sku, 0, supply.low_stock_threshold, supply.max_stock_threshold, 'Out of Stock'])
    return response


def _export_group_csv(queryset, group_name, filename_suffix):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="inventory_{filename_suffix}.csv"'
    writer = csv.writer(response)
    writer.writerow(['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status'])
    for supply in queryset:
        writer.writerow([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, supply.max_stock_threshold, group_name])
    return response


def _export_group_excel(queryset, group_name, filename_suffix):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = group_name
    
    headers = ['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status']
    ws.append(headers)
    
    header_fill = PatternFill(start_color='5ac8fa', end_color='5ac8fa', fill_type='solid')
    header_font = Font(bold=True)
    cell_alignment = Alignment(wrap_text=True)
    
    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).fill = header_fill
        ws.cell(row=1, column=col).font = header_font
    
    row_num = 2
    for supply in queryset:
        ws.append([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, supply.max_stock_threshold, group_name])
        for col in range(1, 7):
            ws.cell(row=row_num, column=col).alignment = cell_alignment
        row_num += 1
    
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col].width = 18
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="inventory_{filename_suffix}.xlsx"'
    return response


def _export_group_pdf(queryset, group_name, filename_suffix):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    
    title = Paragraph(f'<b>{group_name}</b>', styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 20))
    
    cell_style = ParagraphStyle('cell', fontSize=8, wordWrap='LTR')
    
    header_row = [
        Paragraph('Supply Name', cell_style),
        Paragraph('SKU', cell_style),
        Paragraph('Current Stock', cell_style),
        Paragraph('Low Threshold', cell_style),
        Paragraph('Max Threshold', cell_style),
        Paragraph('Status', cell_style),
    ]
    data = [header_row]
    
    for supply in queryset:
        data.append([
            Paragraph(str(supply.name), cell_style),
            Paragraph(str(supply.sku), cell_style),
            Paragraph(str(supply.current_stock), cell_style),
            Paragraph(str(supply.low_stock_threshold), cell_style),
            Paragraph(str(supply.max_stock_threshold), cell_style),
            Paragraph(group_name, cell_style),
        ])
    
    if data:
        col_widths = [100, 60, 60, 60, 60, 60]
        table = Table(data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5ac8fa')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (2, 1), (4, -1), 'CENTER'),
            ('ALIGN', (5, 1), (5, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(table)
    
    doc.build(elements)
    
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="inventory_{filename_suffix}.pdf"'
    return response


def _export_inventory_excel(out_of_stock, low_stock, normal, overstock):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inventory Report'
    
    header_fill = PatternFill(start_color='5ac8fa', end_color='5ac8fa', fill_type='solid')
    header_font = Font(bold=True)
    cell_alignment = Alignment(wrap_text=True)
    group_font = Font(bold=True, size=11)
    group_fill = PatternFill(start_color='e0e0e0', end_color='e0e0e0', fill_type='solid')
    
    def add_group(title, supplies, status):
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = group_font
        ws.cell(row=ws.max_row, column=1).fill = group_fill
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=6)
        
        headers = ['Supply Name', 'SKU', 'Current Stock', 'Low Threshold', 'Max Threshold', 'Status']
        ws.append(headers)
        header_row = ws.max_row
        for col in range(1, 7):
            ws.cell(row=header_row, column=col).fill = header_fill
            ws.cell(row=header_row, column=col).font = header_font
        
        row_num = ws.max_row + 1
        for supply in supplies:
            ws.append([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, supply.max_stock_threshold, status])
            for col in range(1, 7):
                ws.cell(row=row_num, column=col).alignment = cell_alignment
            row_num += 1
    
    add_group('LOW STOCK', low_stock, 'Low Stock')
    add_group('NORMAL STOCK', normal, 'Normal')
    add_group('OVERSTOCK', overstock, 'Overstock')
    add_group('OUT OF STOCK', list(out_of_stock), 'Out of Stock')
    
    for col in ['A', 'B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col].width = 18
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventory_report.xlsx"'
    return response


def _export_inventory_pdf(out_of_stock, low_stock, normal, overstock):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    
    title = Paragraph('<b>Inventory Report</b>', styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 20))
    
    cell_style = ParagraphStyle('cell', fontSize=8, wordWrap='LTR')
    
    out_of_stock_list = list(out_of_stock)
    low_stock_list = list(low_stock)
    normal_list = list(normal)
    overstock_list = list(overstock)
    
    def add_group(title, supplies):
        group_style = ParagraphStyle('group', fontSize=10, fontName='Helvetica-Bold', textColor=colors.HexColor('#333333'))
        elements.append(Paragraph(title, group_style))
        elements.append(Spacer(1, 6))
        
        header_row = [
            Paragraph('Supply Name', cell_style),
            Paragraph('SKU', cell_style),
            Paragraph('Current Stock', cell_style),
            Paragraph('Low Threshold', cell_style),
            Paragraph('Max Threshold', cell_style),
            Paragraph('Status', cell_style),
        ]
        data = [header_row]
        
        for supply in supplies:
            data.append([
                Paragraph(str(supply.name), cell_style),
                Paragraph(str(supply.sku), cell_style),
                Paragraph(str(supply.current_stock), cell_style),
                Paragraph(str(supply.low_stock_threshold), cell_style),
                Paragraph(str(supply.max_stock_threshold), cell_style),
                Paragraph(title.replace('== ', '').replace(' ==', ''), cell_style),
            ])
        
        if supplies:
            col_widths = [100, 60, 60, 60, 60, 60]
            table = Table(data, colWidths=col_widths)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5ac8fa')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (2, 1), (4, -1), 'CENTER'),
                ('ALIGN', (5, 1), (5, -1), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]))
            elements.append(table)
        elements.append(Spacer(1, 15))
    
    add_group('LOW STOCK', low_stock_list)
    add_group('NORMAL STOCK', normal_list)
    add_group('OVERSTOCK', overstock_list)
    add_group('OUT OF STOCK', out_of_stock_list)
    
    doc.build(elements)
    
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="inventory_report.pdf"'
    return response


@login_required
def replenishment_report(request):
    supplies = Supply.objects.filter(is_active=True)
    
    out_of_stock = supplies.filter(current_stock=0).order_by('name')
    low_stock = supplies.filter(
        current_stock__gt=0,
        current_stock__lte=models.F('low_stock_threshold')
    ).order_by('name')

    fmt = request.GET.get('format')
    if fmt == 'csv':
        return _export_replenishment_csv(out_of_stock, low_stock)
    elif fmt == 'excel':
        return _export_replenishment_excel(out_of_stock, low_stock)
    elif fmt == 'pdf':
        return _export_replenishment_pdf(out_of_stock, low_stock)

    context = {
        'out_of_stock': out_of_stock,
        'low_stock': low_stock,
    }
    return render(request, 'inventory/replenishment_report.html', context)


def _get_supplier_names(supply):
    return ', '.join(s.name for s in supply.suppliers.all())


def _export_replenishment_csv(out_of_stock, low_stock):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="replenishment_report.csv"'
    writer = csv.writer(response)
    writer.writerow(['Supply Name', 'SKU', 'Current Stock', 'Suggested Qty', 'Suppliers', 'Status'])
    for supply in out_of_stock:
        writer.writerow([supply.name, supply.sku, 0, supply.low_stock_threshold, _get_supplier_names(supply), 'Out of Stock'])
    for supply in low_stock:
        writer.writerow([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, _get_supplier_names(supply), 'Low Stock'])
    return response


def _export_replenishment_excel(out_of_stock, low_stock):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Replenishment Report'
    
    headers = ['Supply Name', 'SKU', 'Current Stock', 'Suggested Qty', 'Suppliers', 'Status']
    ws.append(headers)
    
    header_fill = PatternFill(start_color='5ac8fa', end_color='5ac8fa', fill_type='solid')
    header_font = Font(bold=True)
    cell_alignment = Alignment(wrap_text=True)
    
    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).fill = header_fill
        ws.cell(row=1, column=col).font = header_font
    
    row_num = 2
    for supply in out_of_stock:
        ws.append([supply.name, supply.sku, 0, supply.low_stock_threshold, _get_supplier_names(supply), 'Out of Stock'])
        for col in range(1, 7):
            ws.cell(row=row_num, column=col).alignment = cell_alignment
        row_num += 1
    for supply in low_stock:
        ws.append([supply.name, supply.sku, supply.current_stock, supply.low_stock_threshold, _get_supplier_names(supply), 'Low Stock'])
        for col in range(1, 7):
            ws.cell(row=row_num, column=col).alignment = cell_alignment
        row_num += 1
    
    for col in ['A', 'B', 'C', 'D', 'F']:
        ws.column_dimensions[col].width = 20
    ws.column_dimensions['E'].width = 30
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="replenishment_report.xlsx"'
    return response


def _export_replenishment_pdf(out_of_stock, low_stock):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    
    title = Paragraph('<b>Replenishment Report</b>', styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 20))
    
    cell_style = ParagraphStyle('cell', fontSize=8, wordWrap='LTR')
    
    header_row = [
        Paragraph('Supply Name', cell_style),
        Paragraph('SKU', cell_style),
        Paragraph('Current Stock', cell_style),
        Paragraph('Suggested Qty', cell_style),
        Paragraph('Suppliers', cell_style),
        Paragraph('Status', cell_style),
    ]
    data = [header_row]
    
    for supply in out_of_stock:
        data.append([
            Paragraph(str(supply.name), cell_style),
            Paragraph(str(supply.sku), cell_style),
            Paragraph('0', cell_style),
            Paragraph(str(supply.low_stock_threshold), cell_style),
            Paragraph(_get_supplier_names(supply), cell_style),
            Paragraph('Out of Stock', cell_style),
        ])
    for supply in low_stock:
        data.append([
            Paragraph(str(supply.name), cell_style),
            Paragraph(str(supply.sku), cell_style),
            Paragraph(str(supply.current_stock), cell_style),
            Paragraph(str(supply.low_stock_threshold), cell_style),
            Paragraph(_get_supplier_names(supply), cell_style),
            Paragraph('Low Stock', cell_style),
        ])
    
    col_widths = [100, 60, 60, 60, 120, 60]
    table = Table(data, colWidths=col_widths)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5ac8fa')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (2, 1), (3, -1), 'CENTER'),
        ('ALIGN', (5, 1), (5, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    
    elements.append(table)
    doc.build(elements)
    
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="replenishment_report.pdf"'
    return response


@login_required
def consumption_report(request):
    months = int(request.GET.get('months', 12))
    custodian_id = request.GET.get('custodian')

    start_date = timezone.now() - timedelta(days=months * 30)
    queryset = SupplyInstallation.objects.filter(
        disposed_at__gte=start_date,
        is_active=False
    )

    if custodian_id:
        queryset = queryset.filter(printer__custodian_id=custodian_id)

    data = queryset.values('disposed_at__year', 'disposed_at__month').annotate(
        count=Count('id')
    ).order_by('disposed_at__year', 'disposed_at__month')

    labels = []
    counts = []
    for item in data:
        labels.append(f"{item['disposed_at__year']}-{item['disposed_at__month']:02d}")
        counts.append(item['count'])

    context = {
        'chart_labels': json.dumps(labels),
        'chart_data': json.dumps(counts),
        'months': months,
        'custodians': Custodian.objects.filter(printers__isnull=False).distinct(),
    }
    return render(request, 'inventory/consumption_report.html', context)


@admin_required
def audit_log(request):
    queryset = AuditLog.objects.all()
    query = request.GET.get('q')
    action = request.GET.get('action')
    user = request.GET.get('user')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    if query:
        queryset = queryset.filter(object_repr__icontains=query)
    if action:
        queryset = queryset.filter(action=action)
    if user:
        queryset = queryset.filter(user_id=user)
    if date_from:
        queryset = queryset.filter(timestamp__date__gte=date_from)
    if date_to:
        queryset = queryset.filter(timestamp__date__lte=date_to)

    queryset = queryset.order_by('-timestamp')
    paginator = Paginator(queryset, 50)
    page = request.GET.get('page')
    logs = paginator.get_page(page)

    context = {
        'logs': logs,
        'actions': AuditLog.ACTION_CHOICES,
        'users': User.objects.all(),
    }
    return render(request, 'inventory/audit_log.html', context)


@login_required
def api_compatible_supplies(request, printer_id):
    printer = get_object_or_404(Printer, pk=printer_id)
    compatible = Supply.objects.filter(
        printer_models=printer.printer_model,
        is_active=True
    ).values('id', 'name', 'sku', 'current_stock')

    supplies = []
    for s in compatible:
        supplies.append({
            'id': s['id'],
            'name': s['name'],
            'sku': s['sku'],
            'current_stock': s['current_stock']
        })

    return JsonResponse({'supplies': supplies})


class SupplyTypeListView(LoginRequiredMixin, ListView):
    model = SupplyType
    template_name = 'inventory/supply_type_list.html'
    paginate_by = 20
    context_object_name = 'supply_types'

    def get_queryset(self):
        queryset = SupplyType.objects.filter(is_active=True).order_by('name')
        query = self.request.GET.get('q')
        if query:
            queryset = queryset.filter(name__icontains=query)
        return queryset


class SupplyTypeCreateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request):
        form = SupplyTypeForm()
        if request.GET.get('modal'):
            return render(request, 'inventory/supply_type_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/supply_type_form.html', {'form': form})

    def post(self, request):
        form = SupplyTypeForm(request.POST)
        if form.is_valid():
            supply_type = form.save()
            messages.success(request, 'Supply type created successfully.')
            
            if request.POST.get('modal'):
                next_url = request.POST.get('next', reverse('supply_type_list'))
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'pk': supply_type.pk, 'name': supply_type.name, 'next': next_url})
                return redirect(next_url + '?created=' + str(supply_type.pk))
            return redirect('supply_type_list')
        if request.POST.get('modal'):
            return render(request, 'inventory/supply_type_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/supply_type_form.html', {'form': form})


class SupplyTypeUpdateView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        supply_type = get_object_or_404(SupplyType, pk=pk)
        form = SupplyTypeForm(instance=supply_type)
        return render(request, 'inventory/supply_type_form.html', {'form': form, 'object': supply_type})

    def post(self, request, pk):
        supply_type = get_object_or_404(SupplyType, pk=pk)
        form = SupplyTypeForm(request.POST, instance=supply_type)
        if form.is_valid():
            form.save()
            messages.success(request, 'Supply type updated successfully.')
            return redirect('supply_type_list')
        return render(request, 'inventory/supply_type_form.html', {'form': form, 'object': supply_type})


class SupplyTypeDeleteView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, pk):
        supply_type = get_object_or_404(SupplyType, pk=pk)
        return render(request, 'inventory/supply_type_confirm_delete.html', {'supply_type': supply_type})

    def post(self, request, pk):
        supply_type = get_object_or_404(SupplyType, pk=pk)
        supply_type.is_active = False
        supply_type.save()
        messages.success(request, 'Supply type deleted successfully.')
        return redirect('supply_type_list')


class SupplyTypeDetailView(LoginRequiredMixin, DetailView):
    model = SupplyType
    template_name = 'inventory/supply_type_detail.html'
    context_object_name = 'supply_type'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        supply_type = self.object
        context['supplies'] = Supply.objects.filter(supply_type=supply_type, is_active=True)
        return context


@login_required
def global_search(request):
    query = request.GET.get('q', '').strip()
    results = {
        'supplies': [],
        'printers': [],
        'suppliers': [],
        'printer_models': [],
    }
    
    if query:
        results['supplies'] = Supply.objects.filter(
            Q(name__icontains=query) | Q(sku__icontains=query),
            is_active=True
        )[:20]
        
        results['printers'] = Printer.objects.filter(
            Q(name__icontains=query) | Q(serial_number__icontains=query)
        ).select_related('printer_model')[:20]
        
        results['suppliers'] = Supplier.objects.filter(
            Q(name__icontains=query) | Q(contact_name__icontains=query),
            is_active=True
        )[:20]
        
        results['printer_models'] = PrinterModel.objects.filter(
            Q(name__icontains=query) | Q(manufacturer__icontains=query)
        )[:20]
    
    return render(request, 'inventory/search_results.html', {
        'query': query,
        'results': results,
    })


@admin_required
def user_list(request):
    users = User.objects.all().order_by('username')
    query = request.GET.get('q')
    role = request.GET.get('role')
    
    if query:
        users = users.filter(
            Q(username__icontains=query) | 
            Q(first_name__icontains=query) | 
            Q(last_name__icontains=query) |
            Q(email__icontains=query)
        )
    if role:
        users = users.filter(role=role)
    
    paginator = Paginator(users, 20)
    page = request.GET.get('page')
    users_page = paginator.get_page(page)
    
    return render(request, 'inventory/user_list.html', {
        'users': users_page,
        'roles': User.ROLE_CHOICES,
    })



@admin_required
def user_create(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.is_superuser = form.cleaned_data.get('is_superuser', False)
            user.is_active = form.cleaned_data.get('is_active', True)
            user.save()
            
            messages.success(request, f'User {user.username} created successfully.')
            log_audit(request.user, 'create', user, {'action': 'created', 'username': user.username})
            return redirect('user_list')
    else:
        form = UserCreationForm()
    
    return render(request, 'inventory/user_form.html', {'form': form})



@admin_required
def user_edit(request, pk):
    user = get_object_or_404(User, pk=pk)
    
    if request.method == 'POST':
        form = UserCreationForm(request.POST, instance=user)
        if form.is_valid():
            user = form.save(commit=False)
            user.is_superuser = form.cleaned_data.get('is_superuser', False)
            user.is_active = form.cleaned_data.get('is_active', True)
            user.save()
            
            messages.success(request, f'User {user.username} updated successfully.')
            log_audit(request.user, 'update', user, {'action': 'updated', 'username': user.username})
            return redirect('user_list')
    else:
        initial = {
            'username': user.username,
            'email': user.email,
            'first_name': user.first_name,
            'last_name': user.last_name,
            'role': user.role,
            'phone': user.phone,
            'is_superuser': user.is_superuser,
            'is_active': user.is_active,
        }
        form = UserCreationForm(initial=initial, instance=user)
    
    return render(request, 'inventory/user_form.html', {'form': form, 'user_obj': user})


@admin_required
def user_delete(request, pk):
    user = get_object_or_404(User, pk=pk)
    
    if request.user.pk == user.pk:
        messages.error(request, 'You cannot delete your own account.')
        return redirect('user_list')
    
    if request.method == 'POST':
        username = user.username
        user.delete()
        messages.success(request, f'User {username} deleted successfully.')
        log_audit(request.user, 'delete', user, {'action': 'deleted', 'username': username})
        return redirect('user_list')
    
    return render(request, 'inventory/user_confirm_delete.html', {'user_obj': user})


@admin_required
def user_toggle_active(request, pk):
    user = get_object_or_404(User, pk=pk)
    
    if request.user.pk == user.pk:
        messages.error(request, 'You cannot toggle your own account status.')
        return redirect('user_list')
    
    user.is_active = not user.is_active
    user.save()
    
    status = "activated" if user.is_active else "deactivated"
    messages.success(request, f'User {user.username} {status}.')
    return redirect('user_list')


@admin_required
def user_reset_password(request, pk):
    user = get_object_or_404(User, pk=pk)
    
    if request.method == 'POST':
        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')
        
        if not new_password or not confirm_password:
            messages.error(request, 'Both password fields are required.')
            return redirect('user_reset_password', pk=pk)
        
        if new_password != confirm_password:
            messages.error(request, 'Passwords do not match.')
            return redirect('user_reset_password', pk=pk)
        
        if len(new_password) < 8:
            messages.error(request, 'Password must be at least 8 characters.')
            return redirect('user_reset_password', pk=pk)
        
        user.set_password(new_password)
        user.save()
        
        messages.success(request, f'Password for {user.username} has been reset.')
        log_audit(request.user, 'update', user, {'action': 'password_reset', 'username': user.username})
        return redirect('user_list')
    
    return render(request, 'inventory/user_reset_password.html', {'user_obj': user})


@admin_required
def database_backup(request):
    from django.conf import settings
    import os
    import shutil
    from datetime import datetime

    backup_dir = os.path.join(settings.BASE_DIR, 'backups')
    os.makedirs(backup_dir, exist_ok=True)

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'backup':
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            db_path = settings.DATABASES['default']['NAME']
            backup_path = os.path.join(backup_dir, f'db_backup_{timestamp}.sqlite3')
            shutil.copy2(db_path, backup_path)
            messages.success(request, f'Database backed up successfully: db_backup_{timestamp}.sqlite3')
            log_audit(request.user, 'backup', None, {'action': 'backup', 'file': f'db_backup_{timestamp}.sqlite3'})

        elif action == 'upload':
            uploaded_file = request.FILES.get('backup_file')
            if uploaded_file:
                if not uploaded_file.name.endswith('.sqlite3'):
                    messages.error(request, 'Invalid file type. Please upload a .sqlite3 file.')
                else:
                    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                    temp_path = os.path.join(backup_dir, f'uploaded_{timestamp}.sqlite3')
                    with open(temp_path, 'wb+') as dest:
                        for chunk in uploaded_file.chunks():
                            dest.write(chunk)
                    messages.success(request, f'Backup uploaded and saved as: uploaded_{timestamp}.sqlite3')
                    log_audit(request.user, 'backup', None, {'action': 'upload', 'file': f'uploaded_{timestamp}.sqlite3'})
            else:
                messages.error(request, 'No file uploaded.')

        elif action == 'restore':
            backup_file = request.POST.get('backup_file')
            if backup_file:
                restore_path = os.path.join(backup_dir, backup_file)
                if os.path.exists(restore_path):
                    db_path = settings.DATABASES['default']['NAME']
                    shutil.copy2(restore_path, db_path)
                    messages.success(request, f'Database restored from: {backup_file}')
                    log_audit(request.user, 'backup', None, {'action': 'restore', 'file': backup_file})
                else:
                    messages.error(request, 'Backup file not found.')

        elif action == 'download':
            backup_file = request.POST.get('backup_file')
            if backup_file:
                restore_path = os.path.join(backup_dir, backup_file)
                if os.path.exists(restore_path):
                    with open(restore_path, 'rb') as f:
                        response = HttpResponse(f.read(), content_type='application/octet-stream')
                        response['Content-Disposition'] = f'attachment; filename="{backup_file}"'
                        return response
                else:
                    messages.error(request, 'Backup file not found.')

        elif action == 'delete':
            backup_file = request.POST.get('backup_file')
            if backup_file:
                delete_path = os.path.join(backup_dir, backup_file)
                if os.path.exists(delete_path):
                    os.remove(delete_path)
                    messages.success(request, f'Backup deleted: {backup_file}')
                    log_audit(request.user, 'backup', None, {'action': 'delete', 'file': backup_file})
                else:
                    messages.error(request, 'Backup file not found.')

        return redirect('database_backup')

    backups = []
    if os.path.exists(backup_dir):
        for f in os.listdir(backup_dir):
            if f.endswith('.sqlite3'):
                full_path = os.path.join(backup_dir, f)
                size = os.path.getsize(full_path)
                mod_time = datetime.fromtimestamp(os.path.getmtime(full_path))
                backups.append({
                    'name': f,
                    'size': size,
                    'modified': mod_time,
                })
    backups.sort(key=lambda x: x['modified'], reverse=True)

    context = {
        'backups': backups,
    }
    return render(request, 'inventory/database_backup.html', context)

# Department Views
class DepartmentListView(LoginRequiredMixin, ListView):
    model = Department
    template_name = 'inventory/department_list.html'
    paginate_by = 20
    context_object_name = 'departments'

    def get_queryset(self):
        queryset = Department.objects.all().order_by('name')
        query = self.request.GET.get('q')
        status = self.request.GET.get('status')
        
        if status != 'all':
            queryset = queryset.filter(is_active=True)
        
        if query:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(description__icontains=query)
            )
        return queryset.prefetch_related('printers', 'custodians')


class DepartmentDetailView(LoginRequiredMixin, DetailView):
    model = Department
    template_name = 'inventory/department_detail.html'
    context_object_name = 'department'


class DepartmentCreateView(LoginRequiredMixin, View):
    def get(self, request):
        form = DepartmentForm()
        if request.GET.get('modal'):
            return render(request, 'inventory/department_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/department_form.html', {'form': form})

    def post(self, request):
        form = DepartmentForm(request.POST)
        if form.is_valid():
            department = form.save()
            messages.success(request, 'Department created successfully.')
            log_audit(request.user, 'create', department, {'action': 'created'})
            
            if request.POST.get('modal'):
                next_url = request.POST.get('next', reverse('department_list'))
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'pk': department.pk, 'name': department.name, 'next': next_url})
                return redirect(next_url + '?created=' + str(department.pk))
            return redirect('department_list')
        if request.POST.get('modal'):
            return render(request, 'inventory/department_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/department_form.html', {'form': form})


class DepartmentUpdateView(LoginRequiredMixin, View):
    def get(self, request, pk):
        department = get_object_or_404(Department, pk=pk)
        form = DepartmentForm(instance=department)
        return render(request, 'inventory/department_form.html', {'form': form, 'object': department})

    def post(self, request, pk):
        department = get_object_or_404(Department, pk=pk)
        form = DepartmentForm(request.POST, instance=department)
        if form.is_valid():
            department = form.save()
            messages.success(request, 'Department updated successfully.')
            log_audit(request.user, 'update', department, {'action': 'updated'})
            return redirect('department_list')
        return render(request, 'inventory/department_form.html', {'form': form, 'object': department})


class DepartmentDeleteView(LoginRequiredMixin, View):
    def get(self, request, pk):
        department = get_object_or_404(Department, pk=pk)
        return render(request, 'inventory/department_confirm_delete.html', {'department': department})

    def post(self, request, pk):
        department = get_object_or_404(Department, pk=pk)
        department.delete()
        messages.success(request, 'Department deleted successfully.')
        log_audit(request.user, 'delete', department, {'action': 'deleted'})
        return redirect('department_list')


class DepartmentToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        department = get_object_or_404(Department, pk=pk)
        department.is_active = not department.is_active
        department.save()
        status = 'enabled' if department.is_active else 'disabled'
        messages.success(request, f'Department {status} successfully.')
        return redirect('department_list')


# Location Views
class LocationListView(LoginRequiredMixin, ListView):
    model = Location
    template_name = 'inventory/location_list.html'
    paginate_by = 20
    context_object_name = 'locations'

    def get_queryset(self):
        queryset = Location.objects.all().order_by('name')
        query = self.request.GET.get('q')
        status = self.request.GET.get('status')
        
        if status != 'all':
            queryset = queryset.filter(is_active=True)
        
        if query:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(description__icontains=query)
            )
        return queryset.prefetch_related('printers', 'custodians')


class LocationDetailView(LoginRequiredMixin, DetailView):
    model = Location
    template_name = 'inventory/location_detail.html'
    context_object_name = 'location'


class LocationCreateView(LoginRequiredMixin, View):
    def get(self, request):
        form = LocationForm()
        if request.GET.get('modal'):
            return render(request, 'inventory/location_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/location_form.html', {'form': form})

    def post(self, request):
        form = LocationForm(request.POST)
        if form.is_valid():
            location = form.save()
            messages.success(request, 'Location created successfully.')
            log_audit(request.user, 'create', location, {'action': 'created'})
            
            if request.POST.get('modal'):
                next_url = request.POST.get('next', reverse('location_list'))
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'pk': location.pk, 'name': location.name, 'next': next_url})
                return redirect(next_url + '?created=' + str(location.pk))
            return redirect('location_list')
        if request.POST.get('modal'):
            return render(request, 'inventory/location_form.html', {'form': form, 'modal': True})
        return render(request, 'inventory/location_form.html', {'form': form})


class LocationUpdateView(LoginRequiredMixin, View):
    def get(self, request, pk):
        location = get_object_or_404(Location, pk=pk)
        form = LocationForm(instance=location)
        return render(request, 'inventory/location_form.html', {'form': form, 'object': location})

    def post(self, request, pk):
        location = get_object_or_404(Location, pk=pk)
        form = LocationForm(request.POST, instance=location)
        if form.is_valid():
            location = form.save()
            messages.success(request, 'Location updated successfully.')
            log_audit(request.user, 'update', location, {'action': 'updated'})
            return redirect('location_list')
        return render(request, 'inventory/location_form.html', {'form': form, 'object': location})


class LocationDeleteView(LoginRequiredMixin, View):
    def get(self, request, pk):
        location = get_object_or_404(Location, pk=pk)
        return render(request, 'inventory/location_confirm_delete.html', {'location': location})

    def post(self, request, pk):
        location = get_object_or_404(Location, pk=pk)
        location.delete()
        messages.success(request, 'Location deleted successfully.')
        log_audit(request.user, 'delete', location, {'action': 'deleted'})
        return redirect('location_list')


class LocationToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        location = get_object_or_404(Location, pk=pk)
        location.is_active = not location.is_active
        location.save()
        status = 'enabled' if location.is_active else 'disabled'
        messages.success(request, f'Location {status} successfully.')
        return redirect('location_list')


# Custodian Views
class CustodianListView(LoginRequiredMixin, ListView):
    model = Custodian
    template_name = 'inventory/custodian_list.html'
    paginate_by = 20
    context_object_name = 'custodians'

    def get_queryset(self):
        queryset = Custodian.objects.all().order_by('first_name', 'last_name')
        query = self.request.GET.get('q')
        department = self.request.GET.get('department')
        location = self.request.GET.get('location')
        status = self.request.GET.get('status')
        
        if status != 'all':
            queryset = queryset.filter(is_active=True)
        
        if query:
            queryset = queryset.filter(
                Q(first_name__icontains=query) | Q(last_name__icontains=query) | Q(email__icontains=query)
            )
        if department:
            queryset = queryset.filter(department_id=department)
        if location:
            queryset = queryset.filter(location_id=location)
        
        return queryset.select_related('department', 'location')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['departments'] = Department.objects.all()
        context['locations'] = Location.objects.all()
        return context


class CustodianCreateView(LoginRequiredMixin, View):
    def get(self, request):
        form = CustodianForm()
        context = {'form': form, 'departments': Department.objects.all(), 'locations': Location.objects.all()}
        if request.GET.get('modal'):
            context['modal'] = True
            return render(request, 'inventory/custodian_form.html', context)
        return render(request, 'inventory/custodian_form.html', context)

    def post(self, request):
        form = CustodianForm(request.POST)
        context = {'form': form, 'departments': Department.objects.all(), 'locations': Location.objects.all()}
        if form.is_valid():
            custodian = form.save()
            messages.success(request, 'Custodian created successfully.')
            log_audit(request.user, 'create', custodian, {'action': 'created'})

            if request.POST.get('modal'):
                next_url = request.POST.get('next', reverse('custodian_list'))
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'pk': custodian.pk, 'name': str(custodian), 'next': next_url})
                return redirect(next_url + '?created=' + str(custodian.pk))
            return redirect('custodian_list')
        if request.POST.get('modal'):
            context['modal'] = True
        return render(request, 'inventory/custodian_form.html', context)


class CustodianUpdateView(LoginRequiredMixin, View):
    def get(self, request, pk):
        custodian = get_object_or_404(Custodian, pk=pk)
        form = CustodianForm(instance=custodian)
        return render(request, 'inventory/custodian_form.html', {'form': form, 'object': custodian, 'departments': Department.objects.all(), 'locations': Location.objects.all()})

    def post(self, request, pk):
        custodian = get_object_or_404(Custodian, pk=pk)
        form = CustodianForm(request.POST, instance=custodian)
        context = {'form': form, 'object': custodian, 'departments': Department.objects.all(), 'locations': Location.objects.all()}
        if form.is_valid():
            custodian = form.save()
            messages.success(request, 'Custodian updated successfully.')
            log_audit(request.user, 'update', custodian, {'action': 'updated'})
            return redirect('custodian_list')
        return render(request, 'inventory/custodian_form.html', context)


class CustodianDeleteView(LoginRequiredMixin, View):
    def get(self, request, pk):
        custodian = get_object_or_404(Custodian, pk=pk)
        return render(request, 'inventory/custodian_confirm_delete.html', {'custodian': custodian})

    def post(self, request, pk):
        custodian = get_object_or_404(Custodian, pk=pk)
        custodian.delete()
        messages.success(request, 'Custodian deleted successfully.')
        log_audit(request.user, 'delete', custodian, {'action': 'deleted'})
        return redirect('custodian_list')


class CustodianToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        custodian = get_object_or_404(Custodian, pk=pk)
        custodian.is_active = not custodian.is_active
        custodian.save()
        status = 'enabled' if custodian.is_active else 'disabled'
        messages.success(request, f'Custodian {status} successfully.')
        return redirect('custodian_list')


class CustodianDetailView(LoginRequiredMixin, DetailView):
    model = Custodian
    template_name = 'inventory/custodian_detail.html'
    context_object_name = 'custodian'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        custodian = self.object
        context['printers'] = Printer.objects.filter(custodian=custodian)
        return context

@login_required
@csrf_exempt
def api_quick_add(request, model_name):
    from django.http import JsonResponse
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Invalid method'})
    
    data = json.loads(request.body)
    
    try:
        if model_name == 'custodian':
            custodian = Custodian.objects.create(
                first_name=data.get('first_name', ''),
                last_name=data.get('last_name', ''),
                email=data.get('email', ''),
                phone=data.get('phone', ''),
                department_id=data.get('department') or None,
                location_id=data.get('location') or None,
                is_active=data.get('is_active', 'true').lower() == 'true'
            )
            return JsonResponse({'success': True, 'id': custodian.id, 'name': str(custodian)})
        
        elif model_name == 'department':
            dept, created = Department.objects.get_or_create(
                name=data.get('name', ''),
                defaults={
                    'description': data.get('description', ''),
                    'is_active': data.get('is_active', 'true').lower() == 'true'
                }
            )
            if not created:
                return JsonResponse({'success': False, 'error': 'Department already exists'})
            return JsonResponse({'success': True, 'id': dept.id, 'name': dept.name})
        
        elif model_name == 'location':
            loc, created = Location.objects.get_or_create(
                name=data.get('name', ''),
                defaults={
                    'description': data.get('description', ''),
                    'is_active': data.get('is_active', 'true').lower() == 'true'
                }
            )
            if not created:
                return JsonResponse({'success': False, 'error': 'Location already exists'})
            return JsonResponse({'success': True, 'id': loc.id, 'name': loc.name})
        
        elif model_name == 'supplier':
            supplier = Supplier.objects.create(
                name=data.get('name', ''),
                contact_name=data.get('contact_name', ''),
                email=data.get('email', ''),
                phone=data.get('phone', ''),
                address=data.get('address', ''),
                is_active=data.get('is_active', 'true').lower() == 'true'
            )
            return JsonResponse({'success': True, 'id': supplier.id, 'name': supplier.name})
        
        elif model_name == 'supplytype':
            supply_type, created = SupplyType.objects.get_or_create(
                name=data.get('name', ''),
                defaults={
                    'description': data.get('description', ''),
                    'is_active': data.get('is_active', 'true').lower() == 'true'
                }
            )
            if not created:
                return JsonResponse({'success': False, 'error': 'Supply type already exists'})
            return JsonResponse({'success': True, 'id': supply_type.id, 'name': supply_type.name})
        
        elif model_name == 'printermodel':
            model, created = PrinterModel.objects.get_or_create(
                name=data.get('name', ''),
                manufacturer=data.get('manufacturer', ''),
                defaults={
                    'description': data.get('description', ''),
                    'is_active': data.get('is_active', 'true').lower() == 'true'
                }
            )
            if not created:
                return JsonResponse({'success': False, 'error': 'Printer model already exists'})
            return JsonResponse({'success': True, 'id': model.id, 'name': str(model)})
        
        else:
            return JsonResponse({'success': False, 'error': 'Invalid model'})
    
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})
